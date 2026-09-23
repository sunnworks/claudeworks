#!/usr/bin/env python3
"""수어 촬영 대본 JSON 을 촬영팀이 그대로 쓸 수 있는 엑셀로 만든다.

사용법: python3 scripts/build-sign-xlsx.py <입력 json> <출력 xlsx>
"""
import json
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

src = sys.argv[1] if len(sys.argv) > 1 else "scripts/out/sign-script.json"
dst = sys.argv[2] if len(sys.argv) > 2 else "scripts/out/수어촬영대본.xlsx"

rows = json.load(open(src, encoding="utf-8"))

HEAD_FILL = PatternFill("solid", fgColor="14459E")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN = Side(style="thin", color="C9D2E4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CATEGORY_FILL = {
    "문항": "E3ECFF",
    "쉬운설명": "F2F6FF",
    "도움말": "FFF7E8",
    "응답안내": "E6F6EE",
    "선택지": "FFFFFF",
    "표항목": "F7F7F7",
    "화면안내": "FDF2F6",
    "버튼": "F4F4F4",
    "안전안내": "FDEAEF",
}

wb = Workbook()

# ---------- 1. 촬영 대본 ----------
ws = wb.active
ws.title = "촬영대본"
headers = [
    "번호", "촬영ID", "구분", "모듈", "사용 위치", "촬영 문장",
    "글자수", "사용 횟수", "촬영 완료", "농인 검수", "의료 검수", "비고",
]
ws.append(headers)
for index, row in enumerate(rows, start=1):
    ws.append([
        index,
        row["id"],
        row["category"],
        row["module"],
        ", ".join(row["usedIn"]),
        row["text"],
        len(row["text"]),
        len(row["usedIn"]),
        "", "", "", "",
    ])

widths = [6, 13, 10, 20, 34, 72, 8, 9, 10, 10, 10, 20]
for col, width in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(col)].width = width

for cell in ws[1]:
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER
ws.row_dimensions[1].height = 30

for excel_row in range(2, ws.max_row + 1):
    category = ws.cell(row=excel_row, column=3).value
    fill = PatternFill("solid", fgColor=CATEGORY_FILL.get(category, "FFFFFF"))
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=excel_row, column=col)
        cell.border = BORDER
        cell.fill = fill
        if col == 6:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        elif col == 5:
            cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="left")
        else:
            cell.alignment = Alignment(vertical="center", horizontal="center")

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

# ---------- 2. 구분별 요약 ----------
summary = wb.create_sheet("구분별 요약")
summary.append(["구분", "문장 수", "설명"])
notes = {
    "문항": "공식 문진표 문구. 한 글자도 바꾸지 않고 그대로 촬영합니다.",
    "쉬운설명": "문항을 쉬운 말로 바꾼 문장. 검증형 척도에는 없습니다.",
    "도움말": "기간·단위·용어를 풀어 주는 문장.",
    "응답안내": "어떻게 고르고 적는지 알려 주는 문장.",
    "선택지": "고르는 항목의 단어. 짧게 촬영합니다.",
    "표항목": "표의 줄 이름(질환명·검사명·술 종류 등).",
    "화면안내": "화면마다 나오는 안내 문장.",
    "버튼": "누르는 버튼 이름.",
    "안전안내": "정신건강 위기 안내. 문구를 바꾸지 않습니다.",
}
counter = Counter(row["category"] for row in rows)
for category in ["문항", "쉬운설명", "도움말", "응답안내", "선택지", "표항목", "화면안내", "버튼", "안전안내"]:
    summary.append([category, counter.get(category, 0), notes[category]])
summary.append(["합계", len(rows), "중복 문장은 한 번만 촬영합니다."])

for col, width in enumerate([14, 10, 60], start=1):
    summary.column_dimensions[get_column_letter(col)].width = width
for cell in summary[1]:
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
for excel_row in range(2, summary.max_row + 1):
    for col in range(1, 4):
        summary.cell(row=excel_row, column=col).border = BORDER
summary.cell(row=summary.max_row, column=1).font = Font(bold=True)
summary.cell(row=summary.max_row, column=2).font = Font(bold=True)

# ---------- 3. 촬영 안내 ----------
guide = wb.create_sheet("촬영 안내")
guide_rows = [
    ["항목", "내용"],
    ["대상", "농인용 건강검진 수어 사전문진 데모에 나오는 모든 문장"],
    ["문장 수", f"{len(rows)}개 (같은 문장은 한 번만 촬영)"],
    ["촬영 ID 규칙", "Q 문항 · E 쉬운설명 · H 도움말 · A 응답안내 · C 선택지 · R 표항목 · G 화면안내 · B 버튼 · S 안전안내"],
    ["파일 이름", "촬영ID 를 그대로 파일 이름으로 씁니다. 예: KSL-Q-001.mp4"],
    ["영상 규격", "가로 영상, 무음, 상반신이 모두 들어오게 촬영합니다. 손이 화면 밖으로 나가지 않아야 합니다."],
    ["배경", "단색 배경. 옷은 손과 대비되는 색으로 입습니다."],
    ["공식 문구", "구분이 '문항'과 '안전안내'인 문장은 뜻을 바꾸지 않습니다. 쉽게 풀어서 표현하지 않습니다."],
    ["검수 순서", "의료진 의미 확인 → 농인 수어전문가 번역 → 촬영 → 농인 검수 → 의료진 재확인"],
    ["연결 방법", "촬영이 끝난 영상을 public/sign-samples/ 에 넣고 src/data/signAssets.ts 의 APPROVED_SIGN_ASSETS 에 등록합니다."],
]
for row in guide_rows:
    guide.append(row)
guide.column_dimensions["A"].width = 16
guide.column_dimensions["B"].width = 96
for cell in guide[1]:
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
for excel_row in range(2, guide.max_row + 1):
    guide.cell(row=excel_row, column=1).font = Font(bold=True)
    for col in (1, 2):
        cell = guide.cell(row=excel_row, column=col)
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)

wb.save(dst)
print(f"엑셀 저장: {dst} (문장 {len(rows)}개)")
